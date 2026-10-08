#!/usr/bin/env python3
"""Validate the Requirements Studio sample fixtures against the schemas and the reference logic.

Document-led sample (samples/*.json, an as-is from documents):
  schema validity; referential integrity; stored confidence == model; gap fingerprints; sample patch applies (AC-M3-1);
  stories + DoR re-derived == stored (M7/M8); as-is flow == reference renderer (M10).
Client KYC sample (samples/client_kyc/, conversation in as-is mode):
  replay every timeline patch from IR v0 across the as-is and the forked to-be, with schema + integrity on every
  intermediate IR (M0/M11); coverage and slot selection match; verbatim evidence for intake answers and suggestions;
  suggestion decisions, ops and benefit maths (M11); replayed IRs == stored; to-be stories + DoR == stored; as-is and to-be
  flows == reference (M10); Jira, Azure DevOps, CSV, Excel and improvements exports == reference (M9); idea records valid (M12).

Usage: python tools/validate_samples.py     (requires jsonschema, jsonpatch, pyyaml, openpyxl)
"""
from __future__ import annotations

import copy, json, sys
from pathlib import Path

import jsonpatch, jsonschema, yaml

sys.path.insert(0, str(Path(__file__).resolve().parent))
import rs_reference as R  # noqa: E402
import render_flow as RF  # noqa: E402
import render_exports as RX  # noqa: E402
import csv, io
import xml.etree.ElementTree as ET

ROOT = Path(__file__).resolve().parent.parent
S, SAMPLES, KYC = ROOT / "schemas", ROOT / "samples", ROOT / "samples" / "client_kyc"
errors: list[str] = []
_validators: dict[str, jsonschema.Draft202012Validator] = {}


def load(p: Path):
    return json.loads(p.read_text())


def check_schema(instance, schema_file: str, label: str):
    if schema_file not in _validators:
        _validators[schema_file] = jsonschema.Draft202012Validator(load(S / schema_file), format_checker=jsonschema.FormatChecker())
    for e in sorted(_validators[schema_file].iter_errors(instance), key=lambda e: list(e.path)):
        errors.append(f"[schema] {label}: /{'/'.join(map(str, e.path))}: {e.message[:200]}")


def exists(ir, i: str) -> bool:
    c = R.PREFIX.get(i.split("_")[0])
    return bool(c) and i in ir.get(c, {}) and ir[c][i].get("meta", {}).get("status") not in R.DEAD


def integrity(ir, label: str):
    def need(i, where):
        if not exists(ir, i):
            errors.append(f"[integrity] {label}: {where} references missing/rejected id '{i}'")

    for eid, e in ir["edges"].items():
        need(e["from"], eid); need(e["to"], eid)
        cond = e.get("condition")
        if cond and cond.get("rule_id"):
            need(cond["rule_id"], eid)
            if cond.get("outcome") not in ir["decision_rules"].get(cond["rule_id"], {}).get("outcomes", []):
                errors.append(f"[integrity] {label}: {eid} outcome '{cond.get('outcome')}' not in {cond['rule_id']}.outcomes")
    for nid, n in ir["nodes"].items():
        if n.get("actor_id"): need(n["actor_id"], nid)
        for k in ("inputs", "outputs", "rule_ids", "sla_ids", "exception_ids", "goal_ids", "system_ids"):
            for i in n.get(k, []): need(i, f"{nid}.{k}")
        if n.get("hitl", {}).get("actor_id"): need(n["hitl"]["actor_id"], f"{nid}.hitl")
        for s in n.get("system_ids", []):
            if s in ir["actors"] and ir["actors"][s]["kind"] != "system":
                errors.append(f"[integrity] {label}: {nid}.system_ids {s} is not a system actor")
    for rid, r in ir["decision_rules"].items():
        for ref in r.get("inputs", []):
            ent, attr = ref.split(".")
            need(ent, rid)
            if ent in ir["entities"] and attr not in {a["name"] for a in ir["entities"][ent]["attributes"]}:
                errors.append(f"[integrity] {label}: {rid} input {ref} attribute undefined")
    for coll in ("exceptions", "slas", "acceptance_criteria", "nfrs"):
        for xid, x in ir[coll].items():
            for i in x["applies_to"]: need(i, xid)
    for xid, x in ir["exceptions"].items():
        for k in ("target_node_id", "notify_actor_id"):
            if x["handling"].get(k): need(x["handling"][k], xid)
    for xid, x in ir["slas"].items():
        for k in ("target_node_id", "notify_actor_id"):
            if x["breach_action"].get(k): need(x["breach_action"][k], xid)
    for gid, g in ir["goals"].items():
        for p in g.get("persona_ids", []): need(p, gid)
    for pid, p in ir["personas"].items():
        if p.get("actor_id"): need(p["actor_id"], pid)
    for coll in R.ELEMENT_COLLS:
        for xid, x in ir[coll].items():
            for p in x["meta"]["provenance"]:
                if p["source_id"] != "src_inferred" and p["source_id"] not in ir["sources"]:
                    errors.append(f"[integrity] {label}: {xid} provenance cites unknown source {p['source_id']}")
    for sid, s in ir.get("stories", {}).items():
        need(s["as_a"], sid)
        for i in s["node_ids"] + s["ac_ids"] + s["goal_ids"] + s["persona_ids"] + s["nfr_ids"]: need(i, sid)


def check_confidence(ir, label):
    for coll in R.ELEMENT_COLLS:
        for xid, x in ir[coll].items():
            want = R.score_meta(ir, x["meta"])
            if abs(want - x["meta"]["confidence"]) > 0.001:
                errors.append(f"[confidence] {label} {xid}: stored {x['meta']['confidence']} != model {want}")


def check_gaps(ir, gaps, label):
    for g in gaps:
        check_schema(g, "gap.schema.json", f"{label}:{g['gap_id']}")
        if R.gap_fingerprint(g) != g["fingerprint"]:
            errors.append(f"[gap] {label} {g['gap_id']}: fingerprint mismatch")
        for ref in g["target_refs"]:
            parts = ref.split("/")
            if parts[1] not in ir or (len(parts) > 2 and parts[2] not in ir[parts[1]]):
                errors.append(f"[gap] {label} {g['gap_id']}: target {ref} not in IR")


def check_stories_and_dor(ir, gaps, questions, dor, label):
    base = copy.deepcopy(ir); base.pop("stories", None)
    want = R.derive_stories(base, gaps, questions, version=ir["process"]["version"])
    rows = R.evaluate_dor(base, want, gaps, signoffs=set(dor.get("signoffs", [])))
    if want != ir["stories"]:
        diff = [k for k in set(want) | set(ir["stories"]) if want.get(k) != ir["stories"].get(k)]
        errors.append(f"[stories] {label}: stored stories differ from M7 derivation for {sorted(diff)}")
    if rows != dor["stories"]:
        errors.append(f"[dor] {label}: stored DoR report differs from M8 evaluation")


def check_flow(ir, folder: Path, label):
    prefix = ir["process"]["variant"] + "_"
    """M10: stored diagrams == reference render; draw.io parses; every live node drawn with its control style."""
    for fname, fn in ((prefix + "process_flow.drawio", RF.render_drawio), (prefix + "process_flow.mmd", RF.render_mermaid)):
        f = folder / fname
        if not f.exists():
            errors.append(f"[flow] {label}: missing {fname}"); continue
        if f.read_text() != fn(ir):
            errors.append(f"[flow] {label}: {fname} differs from the reference renderer")
    try:
        root = ET.parse(folder / (prefix + "process_flow.drawio")).getroot()
    except Exception as exc:  # noqa: BLE001
        errors.append(f"[flow] {label}: draw.io XML does not parse: {exc}"); return
    cells = {c.get("id"): c for c in root.iter("mxCell")}
    for nid in R.live(ir, "nodes"):
        if nid not in cells:
            errors.append(f"[flow] {label}: node {nid} not drawn"); continue
        want = RF.STYLE.get(RF.control(ir, nid))
        if want and cells[nid].get("style") != want:
            errors.append(f"[flow] {label}: {nid} drawn with the wrong control style")
    for c in cells.values():
        if c.get("edge") == "1" and (c.get("source") not in cells or c.get("target") not in cells):
            errors.append(f"[flow] {label}: edge {c.get('id')} points at a missing shape")
    mmd = (folder / (prefix + "process_flow.mmd")).read_text()
    for nid in R.live(ir, "nodes"):
        if nid not in mmd:
            errors.append(f"[flow] {label}: node {nid} missing from Mermaid")


def document_led():
    ir = load(SAMPLES / "ir_client_onboarding.json")
    gaps = load(SAMPLES / "gaps_client_onboarding.json")
    questions = load(SAMPLES / "questions_outbox.json")
    patch = load(SAMPLES / "patch_example.json")
    dor = load(SAMPLES / "dor_report.json")
    smes = {s["sme_id"] for s in load(SAMPLES / "sme_directory.json")}
    check_schema(ir, "process-ir.schema.json", "ir_client_onboarding")
    for q in questions: check_schema(q, "question.schema.json", q["question_id"])
    check_schema(patch, "patch.schema.json", "patch_example")
    integrity(ir, "ir_client_onboarding")
    check_confidence(ir, "ir_client_onboarding")
    check_gaps(ir, gaps, "client_onboarding")
    gap_ids = {g["gap_id"] for g in gaps}
    for q in questions:
        for gid in q["gap_ids"]:
            if gid not in gap_ids: errors.append(f"[question] {q['question_id']}: unknown gap {gid}")
        if q["sme_id"] not in smes: errors.append(f"[question] {q['question_id']}: unknown SME {q['sme_id']}")
    check_stories_and_dor(ir, gaps, [], dor, "client_onboarding")
    if ir["nodes"]["node_screening"]["meta"]["confidence"] != 0.834:
        errors.append("[confidence] node_screening should be 0.834 (AC-M8-1)")
    # AC-M3-1
    try:
        after = jsonpatch.apply_patch(copy.deepcopy(ir), patch["ops"])
        after.pop("stories", None); R.score_all(after)
        check_schema(after, "process-ir.schema.json", "ir_after_patch"); integrity(after, "ir_after_patch")
        n = after["nodes"]["node_high_risk_approval"]
        assert n["actor_id"] == "act_mlro" and n["meta"]["status"] == "confirmed" and n["meta"]["confidence"] >= 0.95
    except Exception as exc:  # noqa: BLE001
        errors.append(f"[patch] AC-M3-1 failed: {exc}")
    yaml.safe_load((SAMPLES / "export" / "manifest.yaml").read_text())
    check_flow(ir, SAMPLES / "flow", "client_onboarding")
    return f"document-led: {len(ir['nodes'])} nodes, {len(ir['stories'])} stories, {len(gaps)} gaps, coverage {R.coverage(ir)['percent']}"


def strip(ir):
    x = copy.deepcopy(ir); x.pop("stories", None)
    x["process"] = {k: x["process"][k] for k in ("id", "name")}
    for c in R.ELEMENT_COLLS:
        for e in x[c].values(): e["meta"]["confidence"] = 0.0
    return x


def client_kyc():
    """As-is interview -> fork -> AI suggestions -> to-be -> stories, flows, exports (M0, M11, M7, M10, M9)."""
    session = load(KYC / "intake_session_kyc.json")
    asis_final, tobe_final = load(KYC / "ir_client_kyc_as_is.json"), load(KYC / "ir_client_kyc_to_be.json")
    gaps, dor = load(KYC / "gaps_client_kyc.json"), load(KYC / "dor_report_kyc.json")
    suggestions, idea = load(KYC / "suggestions_client_kyc.json"), load(KYC / "idea_client_kyc.json")
    smes = {s["sme_id"] for s in load(SAMPLES / "sme_directory.json")}
    check_schema(session, "intake-session.schema.json", "intake_session")
    check_schema(idea, "idea.schema.json", "idea_client_kyc")
    for i in load(SAMPLES / "ideas_index.json"):
        check_schema(i, "idea.schema.json", f"ideas_index:{i['idea_id']}")
    for x in suggestions:
        check_schema(x, "suggestion.schema.json", x["suggestion_id"])
    for q in session["parked_questions"]:
        check_schema(q, "question.schema.json", q["question_id"])
        if q["sme_id"] not in smes: errors.append(f"[question] {q['question_id']}: unknown SME {q['sme_id']}")
    for lbl, ir_ in (("as_is", asis_final), ("to_be", tobe_final)):
        check_schema(ir_, "process-ir.schema.json", f"ir_{lbl}"); integrity(ir_, f"ir_{lbl}"); check_confidence(ir_, f"ir_{lbl}")
    pid = {"as_is": asis_final["process"]["id"], "to_be": tobe_final["process"]["id"]}
    for v, ir_ in (("as_is", asis_final), ("to_be", tobe_final)):
        check_gaps(ir_, [g for g in gaps if g["process_id"] == pid[v]], f"client_kyc_{v}")

    irs = {"as_is": load(KYC / "ir_v0_empty.json"), "to_be": None}
    check_schema(irs["as_is"], "process-ir.schema.json", "ir_v0")
    answers, slot_turns, applied, decided = {}, 0, 0, {}
    sug_by = {x["suggestion_id"]: x for x in suggestions}
    for e in session["timeline"]:
        v = e.get("process", "as_is")
        if e["kind"] == "fork":
            f = e["fork"]
            irs["to_be"] = copy.deepcopy(irs["as_is"]); irs["to_be"]["process"].update(f["process_overrides"]); continue
        if e.get("turn"):
            answers[e["turn"]] = e["answer"]["text"]
        if e["kind"] == "turn" and e["target"]["kind"] == "slot":
            slot_turns += 1
            want = R.next_slot(irs[v])
            if want != e["target"]["id"]:
                errors.append(f"[intake] {e['turn']}: target {e['target']['id']} but next_slot() on pre-turn IR = {want}")
        if e["kind"] == "suggestion_decision":
            decided[e["suggestion_id"]] = e["decision"]
            x = sug_by[e["suggestion_id"]]
            if x["status"] != e["decision"]:
                errors.append(f"[improve] {x['suggestion_id']}: status {x['status']} != decision {e['decision']}")
            if e["decision"] == "accepted" and x["ops"] != e.get("ops", []):
                errors.append(f"[improve] {x['suggestion_id']}: suggestion ops differ from the applied patch")
            if e["decision"] == "rejected" and e.get("ops"):
                errors.append(f"[improve] {x['suggestion_id']}: rejected suggestion changed the to-be")
        if e.get("ops"):
            try:
                irs[v] = jsonpatch.apply_patch(irs[v], e["ops"])
            except Exception as exc:  # noqa: BLE001
                errors.append(f"[intake] seq {e['seq']}: patch failed: {exc}"); break
            applied += 1
            lbl = f"replay@{v}:{e.get('turn') or e.get('suggestion_id') or e['kind']}"
            check_schema(irs[v], "process-ir.schema.json", lbl); integrity(irs[v], lbl)
        if "coverage_after" in e:
            got = R.coverage(irs[v])
            if got["percent"] != e["coverage_after"]["percent"] or got["filled"] != e["coverage_after"]["filled"]:
                errors.append(f"[intake] seq {e['seq']}: coverage {got['percent']} != recorded {e['coverage_after']['percent']}")
    if set(decided) != set(sug_by):
        errors.append(f"[improve] undecided suggestions: {sorted(set(sug_by) - set(decided))}")
    sme_answers = {q["answer"]["text"] for q in session["parked_questions"] if q.get("answer")}
    for x in suggestions:
        for evd in x["evidence"]:
            pool = [answers.get(evd["locator"], "")] if evd["source_id"] == session["source_id"] else list(sme_answers)
            if not any(evd["excerpt"] in t for t in pool):
                errors.append(f"[improve] {x['suggestion_id']}: evidence not verbatim: '{evd['excerpt'][:50]}'")
        b = x["benefit"]
        if b.get("minutes_saved_per_case") is not None and b.get("hours_saved_per_month") != round(b["minutes_saved_per_case"] * b["cases_per_month"] / 60, 1):
            errors.append(f"[improve] {x['suggestion_id']}: hours_saved_per_month does not match minutes × volume")
    for ir_ in (asis_final, tobe_final):
        for coll in R.ELEMENT_COLLS:
            for xid, x in ir_[coll].items():
                for p in x["meta"]["provenance"]:
                    if p["source_id"] == session["source_id"] and p["excerpt"] not in answers.get(p["locator"]["value"], ""):
                        errors.append(f"[intake] {xid}: excerpt not verbatim in {p['locator']['value']}: '{p['excerpt'][:60]}'")
                    if p["locator"]["kind"] == "suggestion" and p["excerpt"] != sug_by[p["locator"]["value"]]["title"]:
                        errors.append(f"[improve] {xid}: suggestion provenance excerpt must equal the suggestion title")
    for v, final in (("as_is", asis_final), ("to_be", tobe_final)):
        if strip(irs[v]) != strip(final):
            errors.append(f"[intake] replayed {v} IR differs from the stored file")
    tb_gaps = [g for g in gaps if g["process_id"] == pid["to_be"]]
    check_stories_and_dor(tobe_final, tb_gaps, session["parked_questions"], dor, "client_kyc_to_be")
    check_flow(asis_final, KYC / "flow", "client_kyc_as_is"); check_flow(tobe_final, KYC / "flow", "client_kyc_to_be")
    gates = [k for k, n in tobe_final["nodes"].items() if n.get("hitl", {}).get("mode") == "approval"]
    for k in gates:
        for prob in R.approval_problems(tobe_final, k):
            errors.append(f"[hitl] {prob}")
    if len(gates) < 2: errors.append("[hitl] KYC to-be should keep the analyst and MLRO approval gates")
    changed = [k for k, n in tobe_final["nodes"].items() if n.get("change")]
    acc = {x["suggestion_id"] for x in suggestions if x["status"] == "accepted"}
    if {tobe_final["nodes"][k]["change"]["suggestion_id"] for k in changed} != acc:
        errors.append("[improve] to-be change markers do not match the accepted suggestions")
    not_ready = [r["story_id"] for r in dor["stories"] if r["status"] != "ready"]
    if not_ready: errors.append(f"[intake] stories not ready: {not_ready}")
    # exports (M9)
    X, st, title, desc = KYC / "exports", tobe_final["stories"], idea["title"], tobe_final["process"]["description"]
    for fname, want in (("jira_import.csv", RX.jira_csv(tobe_final, st, title, desc)), ("azure_devops_import.csv", RX.ado_csv(tobe_final, st, title, desc)),
                        ("stories_backlog.csv", RX.backlog_csv(tobe_final, st)), ("improvements.md", RX.improvements_md(asis_final, tobe_final, suggestions))):
        if (X / fname).read_text() != want:
            errors.append(f"[export] {fname} differs from the reference renderer")
    jira = list(csv.reader(io.StringIO((X / "jira_import.csv").read_text())))
    if jira[1][1] != "Epic" or any(r[3] != "1" for r in jira[2:]) or len(jira) - 2 != len(st):
        errors.append("[export] Jira CSV must list the epic first and link every story to it via Parent")
    ado = list(csv.reader(io.StringIO((X / "azure_devops_import.csv").read_text())))
    if ado[1][1] != "Feature" or any(r[1] != "User Story" or r[2] or not r[3] for r in ado[2:]) or any(r[0] for r in ado[1:]):
        errors.append("[export] Azure DevOps CSV must have an empty ID, a Feature in Title 1 and stories in Title 2")
    from openpyxl import load_workbook
    wb = load_workbook(X / "stories_backlog.xlsx")
    want = RX.workbook_rows(tobe_final, st, asis_final, suggestions)
    for name, header in RX.SHEETS.items():
        rows = [list(r) for r in wb[name].iter_rows(values_only=True)]
        norm = lambda r: ["" if c is None else c for c in r]
        if norm(rows[0]) != header or [norm(r) for r in rows[1:]] != [norm(r) for r in want[name]]:
            errors.append(f"[export] stories_backlog.xlsx sheet '{name}' differs from the reference rows")
    hours = sum(x["benefit"]["hours_saved_per_month"] or 0 for x in suggestions if x["status"] == "accepted")
    return (f"client KYC: {applied} patches replayed (as-is + forked to-be), {slot_turns} slot selections verified, "
            f"{len(acc)}/{len(suggestions)} suggestions accepted (~{hours:g} h/month), {len(st)} stories ready, "
            f"2 flows, {len(gates)} approval gates, Jira/ADO/CSV/xlsx exports match")


def main() -> int:
    a = document_led()
    b = client_kyc()
    if errors:
        print(f"FAILED — {len(errors)} problem(s):")
        for e in errors[:60]: print("  " + e)
        return 1
    print("OK — " + a); print("OK — " + b)
    return 0


if __name__ == "__main__":
    sys.exit(main())
