"""Backlog CSV and the Excel workbook (M9). Port of tools/render_exports.py; the CSV is the Stories sheet."""
import io

from flow_renderer import hitl_summary

from .common import MODE, actor, change_line, gwt_lines, to_csv
from .compare import compare_rows

BACKLOG_HEADER = ["Story ID", "Title", "Priority", "Control", "As a", "I want", "So that", "Acceptance criteria",
                  "Edge cases", "Non-functional", "Upstream", "Downstream", "Systems", "Open questions",
                  "Change from today", "DoR", "Confidence"]
SHEETS = {
    "Stories": BACKLOG_HEADER,
    "Acceptance criteria": ["Story ID", "AC ID", "Kind", "Title", "Given", "When", "Then"],
    "Human controls": ["Step", "Control", "Who", "When", "Checks"],
    "Improvements": ["ID", "Suggestion", "Kind", "Decision", "Minutes saved per case", "Hours saved per month",
                     "Human control", "Reason"],
    "As-is vs to-be": ["Step", "Today", "Minutes per case today", "To-be", "Change"],
}


def backlog_rows(ir: dict, stories: dict) -> list[list]:
    rows = []
    for sid, s in stories.items():
        acs = []
        for i, aid in enumerate(s["ac_ids"], 1):
            ac = ir["acceptance_criteria"][aid]
            acs.append(f"AC{i} {ac['title']}: " + " ".join(f"{kw} {line}." for kw, line in gwt_lines(ac)))
        dep = s["dependencies"]
        rows.append([sid, s["title"], s["priority"] or "", MODE[s["human_control"]["mode"]], actor(ir, s["as_a"]),
                     f"to {s['i_want']}", s["so_that"], "\n".join(acs), "\n".join(x["title"] for x in s["edge_cases"]),
                     "\n".join(ir["nfrs"][n]["statement"] for n in s["nfr_ids"]),
                     ", ".join(dep["upstream_story_ids"]), ", ".join(dep["downstream_story_ids"]),
                     ", ".join(actor(ir, a) for a in dep["system_actor_ids"]),
                     "\n".join(q["text"] for q in s["open_questions"]), change_line(s), s["dor_status"],
                     s["confidence"]])
    return rows


def backlog_csv(ir: dict, stories: dict) -> str:
    return to_csv(backlog_rows(ir, stories), BACKLOG_HEADER)


def ac_rows(ir: dict, stories: dict) -> list[list]:
    rows = []
    for sid, s in stories.items():
        for aid in s["ac_ids"]:
            ac = ir["acceptance_criteria"][aid]
            rows.append([sid, aid, ac.get("kind", "happy_path"), ac["title"],
                         *("\n".join(ac[k]) for k in ("given", "when", "then"))])
    return rows


def suggestion_rows(suggestions: list[dict]) -> list[list]:
    return [[x["suggestion_id"], x["title"], x["kind"].replace("_", " "), x["status"],
             x["benefit"].get("minutes_saved_per_case"), x["benefit"].get("hours_saved_per_month"), x["controls"],
             (x.get("decision") or {}).get("reason", "")] for x in suggestions]


def workbook_rows(to_be: dict, stories: dict, as_is: dict | None = None, suggestions=()) -> dict[str, list[list]]:
    humans = [[r["step"], r["control"].replace("_", " "), r["who"] or "", r["when"].replace("_", " "),
               r["criteria"] or ""] for r in hitl_summary(to_be)]
    return {"Stories": backlog_rows(to_be, stories), "Acceptance criteria": ac_rows(to_be, stories),
            "Human controls": humans, "Improvements": suggestion_rows(list(suggestions)),
            "As-is vs to-be": compare_rows(as_is, to_be)}


def backlog_xlsx(to_be: dict, stories: dict, as_is: dict | None = None, suggestions=()) -> bytes:
    """Header row styled and frozen, auto-filter on, wrapped text (M9)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    data = workbook_rows(to_be, stories, as_is, suggestions)
    for name, header in SHEETS.items():
        ws = wb.create_sheet(name)
        ws.append(header)
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="2F4858")
        for row in data[name]:
            ws.append(row)
        for col in range(1, len(header) + 1):
            ws.column_dimensions[get_column_letter(col)].width = 18 if len(header) > 8 else 28
            for column in ws.iter_cols(min_col=col, max_col=col, min_row=2):
                for cell in column:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
    wb.properties.creator = "Requirements Studio"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
