"""M9 exports: reference parity (pure), and the export service on the seeded samples against Postgres."""
import csv
import io
import os
import zipfile

import pytest
from openpyxl import load_workbook

from exporters import SHEETS, ado_csv, backlog_csv, export_hints, jira_csv, workbook_rows
from services.common.settings import Settings
from services.export.service import ExportService
from services.export.store import AzureBlobStore, FileStore
from services.ideas.seed import seed
from story_renderer import derive_stories
from tests.conftest import KYC, load
from tests.oracle import _load, replayed_irs

RX = _load("render_exports")
EXPORTS = KYC / "exports"
ALL = ["lucidchart", "mermaid", "markdown", "gherkin", "jira", "ado", "excel", "csv"]
TO_BE = load(KYC / "ir_client_kyc_to_be.json")


# ---------------------------------------------------------------- pure parity with tools/render_exports.py

IRS = [pytest.param(TO_BE, id="kyc_to_be")] + [pytest.param(ir, id=k) for k, ir in replayed_irs()
                                                 if k.startswith("to_be")]


@pytest.mark.parametrize("ir", IRS)
def test_m9_renderers_match_reference(ir):
    stories = ir.get("stories") or derive_stories(ir, [])
    as_is, suggestions = load(KYC / "ir_client_kyc_as_is.json"), load(KYC / "suggestions_client_kyc.json")
    assert jira_csv(ir, stories, "T", "D") == RX.jira_csv(ir, stories, "T", "D")
    assert ado_csv(ir, stories, "T", "D") == RX.ado_csv(ir, stories, "T", "D")
    assert backlog_csv(ir, stories) == RX.backlog_csv(ir, stories)
    assert workbook_rows(ir, stories, as_is, suggestions) == RX.workbook_rows(ir, stories, as_is, suggestions)


def test_M9_AC_M9_5():
    """qc_rubric_seed for a task equals the then clauses of its acceptance criteria."""
    hints = export_hints(TO_BE)
    for nid, n in TO_BE["nodes"].items():
        if n["type"] != "task":
            continue
        acs = sorted(a for a, ac in TO_BE["acceptance_criteria"].items() if nid in ac["applies_to"])
        assert hints[nid]["acceptance_criteria_ids"] == acs
        assert hints[nid]["qc_rubric_seed"] == [t for a in acs for t in TO_BE["acceptance_criteria"][a]["then"]]
    assert hints["node_record_kyc"]["qc_rubric_seed"][:3] == ["recording is retried up to 3 times",
                                                              "after the third failure the case is placed in the "
                                                              "analyst's queue",
                                                              "nothing is marked KYC complete until the CRM write "
                                                              "succeeds"]
    assert hints["node_screen"]["hitl"]["mode"] == "hitl_review"


def test_m9_ado_scrum_names_product_backlog_items():
    rows = list(csv.reader(io.StringIO(ado_csv(TO_BE, TO_BE["stories"], "T", "D", process="scrum"))))
    assert {r[1] for r in rows[2:]} == {"Product Backlog Item"} and rows[1][1] == "Feature"
    with pytest.raises(ValueError):
        ado_csv(TO_BE, TO_BE["stories"], "T", "D", process="kanban")


# ---------------------------------------------------------------- the service

@pytest.fixture
async def db(tx_sessionmaker):
    async with tx_sessionmaker() as session:
        yield session


@pytest.fixture
def exports(db, tmp_path):
    return ExportService(db, Settings(env="dev"), FileStore(tmp_path))


async def _files(exports, export_id):
    data, _, _ = await exports.download(export_id)
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        return {n: z.read(n) for n in z.namelist()}


async def test_M9_AC_M9_1(db, exports):
    """For the client KYC to-be, Jira, ADO, backlog CSV and improvements equal the samples byte for byte, and the
    .xlsx sheets equal the reference rows."""
    await seed(db, "samples")
    export = await exports.create("idea_client_kyc", ALL, user_id="user_sarah_lin")
    files = await _files(exports, export["export_id"])
    for name in ("jira_import.csv", "azure_devops_import.csv", "stories_backlog.csv", "improvements.md"):
        assert files[name] == (EXPORTS / name).read_bytes(), name
    for name in ("as_is_process_flow.drawio", "to_be_process_flow.drawio", "as_is_process_flow.mmd",
                 "to_be_process_flow.mmd"):
        assert files[f"flow/{name}"] == (KYC / "flow" / name).read_bytes(), name
    got = load_workbook(io.BytesIO(files["stories_backlog.xlsx"]))
    want = load_workbook(EXPORTS / "stories_backlog.xlsx")
    for sheet in SHEETS:
        assert [list(r) for r in got[sheet].iter_rows(values_only=True)] == \
            [list(r) for r in want[sheet].iter_rows(values_only=True)], sheet
    ws = got["Stories"]
    assert ws.freeze_panes == "A2" and ws.auto_filter.ref and ws["A1"].font.bold
    assert "features/client_kyc.feature" in files and "Draft" not in files["stories.md"].decode()
    assert export["draft"] is False and export["ir_version"] == 12


async def test_M9_AC_M9_2(db, exports):
    """Jira: the Epic first, every Story has Parent = 1. ADO: empty ID, the Feature in Title 1, stories in Title 2."""
    await seed(db, "samples")
    export = await exports.create("idea_client_kyc", ["jira", "ado"], user_id="user_sarah_lin")
    files = await _files(exports, export["export_id"])
    jira = list(csv.reader(io.StringIO(files["jira_import.csv"].decode())))
    assert jira[1][:2] == ["1", "Epic"] and len(jira) - 2 == 9
    assert all(r[1] == "Story" and r[3] == "1" for r in jira[2:])
    ado = list(csv.reader(io.StringIO(files["azure_devops_import.csv"].decode())))
    assert ado[0][:4] == ["ID", "Work Item Type", "Title 1", "Title 2"]
    assert all(r[0] == "" for r in ado[1:])
    assert ado[1][1:3] == ["Feature", "Client KYC checks"] and ado[1][3] == ""
    assert all(r[1] == "User Story" and r[2] == "" and r[3] for r in ado[2:]) and len(ado) - 2 == 9


async def test_m9_draft_idea_carries_banner_and_labels(db, exports):
    await seed(db, "samples")
    export = await exports.create("idea_client_onboarding", ["markdown", "jira", "ado"], user_id="user_daniel_okafor")
    files = await _files(exports, export["export_id"])
    assert export["draft"] is True
    assert "> **Draft — 9 stories not ready.**" in files["stories.md"].decode()
    jira = list(csv.reader(io.StringIO(files["jira_import.csv"].decode())))
    assert jira[0].count("Labels") == 4 and all(r[8] == "DRAFT" for r in jira[1:])
    ado = list(csv.reader(io.StringIO(files["azure_devops_import.csv"].decode())))
    assert all(r[7].endswith("; DRAFT") for r in ado[1:])
    assert "improvements.md" not in files  # no as-is → to-be improvement for a document-led idea yet


async def test_m9_same_content_is_the_same_export_and_history_supersedes(db, exports):
    await seed(db, "samples")
    first = await exports.create("idea_client_kyc", ["jira", "excel"], user_id="user_sarah_lin")
    again = await exports.create("idea_client_kyc", ["excel", "jira"], user_id="user_sarah_lin")
    assert again["export_id"] == first["export_id"]
    other = await exports.create("idea_client_kyc", ["csv"], user_id="user_sarah_lin")
    history = await exports.history("idea_client_kyc")
    assert [h["export_id"] for h in history] == [other["export_id"], first["export_id"]]
    assert history[1]["superseded_by"] == other["export_id"]
    assert history[0]["formats"] == ["csv"] and history[0]["created_by"] == "user_sarah_lin"


async def test_m9_single_file_download_and_preview(db, exports):
    await seed(db, "samples")
    export = await exports.create("idea_client_kyc", ["jira", "csv"], user_id="user_sarah_lin")
    content, name, kind = await exports.download(export["export_id"], "jira_import.csv")
    assert content == (EXPORTS / "jira_import.csv").read_bytes() and name == "jira_import.csv" and kind == "text/csv"
    preview = await exports.preview("idea_client_kyc", "csv")
    assert preview["files"][0]["text"].startswith("Story ID,Title,Priority")
    formats = await exports.formats("idea_client_kyc")
    assert {f["id"] for f in formats["formats"] if f["default"]} == {"lucidchart", "markdown", "jira"}
    with pytest.raises(Exception, match="choose formats"):
        await exports.create("idea_client_kyc", ["pdf"], user_id="user_sarah_lin")


AZURITE = os.environ.get("RS_AZURE_STORAGE_CONNECTION_STRING")


@pytest.mark.skipif(not AZURITE, reason="RS_AZURE_STORAGE_CONNECTION_STRING not set (docker compose azurite)")
async def test_m9_azure_blob_store_round_trip():
    store = AzureBlobStore(container="exports-test", connection_string=AZURITE)  # created on first use
    try:
        uri = await store.put("proc_x/roundtrip.zip", b"zip-bytes")
        assert uri.endswith("/exports-test/proc_x/roundtrip.zip")
        assert await store.get("proc_x/roundtrip.zip") == b"zip-bytes"
    finally:
        await store.close()


@pytest.mark.skipif(not AZURITE, reason="RS_AZURE_STORAGE_CONNECTION_STRING not set (docker compose azurite)")
async def test_m9_export_through_azure_blob(db):
    await seed(db, "samples")
    store = AzureBlobStore(container="exports-test", connection_string=AZURITE)
    try:
        exports = ExportService(db, Settings(env="dev"), store)
        export = await exports.create("idea_client_kyc", ["jira"], user_id="user_sarah_lin")
        content, name, _ = await exports.download(export["export_id"], "jira_import.csv")
        assert content == (EXPORTS / "jira_import.csv").read_bytes()
    finally:
        await store.close()
